"""Throwaway: mine learn/*.xlsx for per-discipline title phrases + type purity.

Run: .venv/Scripts/python.exe scripts/mine_dossiers.py
Not imported by the package. Output is hand-authored into config.
"""
from __future__ import annotations

import math
from collections import Counter, defaultdict

import openpyxl

from classifier.core.scoring import canonicalize_title, candidate_phrases

FOLD = {
    'PROJECT MANAGEMENT': 'EMT', 'PROJECT': 'EMT',
    'PROCESS': 'PRO', 'HSE': 'PRO', 'HSE AND PROCESS': 'PRO', 'SAFETY/RISK': 'PRO',
    'ELECTRICAL': 'ELE', 'INSTRUMENTATION': 'INC', 'INSTRUMENT': 'INC',
    'MECHANICAL': 'MEC', 'ROTATING': 'MEC', 'ROTARY': 'MEC',
    'PIPING': 'PIP', 'PIPELINE': 'PIP', 'CIVIL': 'CIV', 'STRUCTURAL': 'CIV',
}
CODE2ID = {'CIV': 1, 'CRM': 2, 'ELE': 3, 'EMT': 4, 'HSE': 5, 'INC': 6, 'MEC': 7,
           'PIP': 8, 'PKG': 9, 'PRJ': 10, 'PRO': 11, 'PRP': 12, 'QAC': 13,
           'RO': 14, 'STC': 36, 'TEL': 37, 'HVAC': 39}


def load_records():
    records = []  # (title, type_or_None, code)
    wb = openpyxl.load_workbook(
        'learn/50703-EPC Dossier Index - PDF (Updated on 27-03-2014).xlsx',
        read_only=True, data_only=True)
    ws = wb['PDF']
    cur = None
    for row in ws.iter_rows(values_only=True):
        c0, title, typ = row[0], row[1], row[4]
        if c0 and not title:
            u = str(c0).strip().upper()
            if u in FOLD:
                cur = FOLD[u]
            continue
        if title and typ and typ != 'Type':
            records.append((str(title), str(typ), cur))
    wb2 = openpyxl.load_workbook('learn/P03433 - EPC  DOSSIER.xlsx',
                                 read_only=True, data_only=True)
    ws2 = wb2['FEED DOSSIER']
    for row in ws2.iter_rows(values_only=True):
        desc, disc = row[7], row[8]
        if desc and disc:
            u = str(disc).strip().upper()
            if u in FOLD:
                records.append((str(desc), None, FOLD[u]))
    return records


def main():
    records = load_records()
    # phrase counts per discipline code
    per_disc = defaultdict(Counter)   # code -> phrase -> count
    disc_titles = Counter()
    for title, _typ, code in records:
        if not code:
            continue
        disc_titles[code] += 1
        toks = canonicalize_title(title)
        seen = set()
        for ph in candidate_phrases(toks):
            p = ' '.join(ph)
            if p in seen:
                continue
            seen.add(p)
            per_disc[code][p] += 1

    # df across disciplines (how many disciplines contain the phrase)
    df = Counter()
    for code, cnt in per_disc.items():
        for p in cnt:
            df[p] += 1
    n_disc = len(per_disc)

    print('=== per-discipline title counts ===')
    for code, n in disc_titles.most_common():
        print(f'  {code} (id {CODE2ID[code]}): {n}')

    print('\n=== top discriminative phrases per discipline ===')
    for code in sorted(per_disc, key=lambda c: -disc_titles[c]):
        cnt = per_disc[code]
        scored = []
        for p, c in cnt.items():
            if c < 2:
                continue
            idf = math.log(n_disc / df[p])      # 0 when phrase in all disciplines
            disc = c * (idf + 0.3)              # favour frequent + concentrated
            # weight in the existing config style: log-scaled, lifted by purity
            weight = round(math.log1p(c) + idf, 4)
            scored.append((disc, weight, c, df[p], p))
        scored.sort(reverse=True)
        print(f'\n-- {code} (id {CODE2ID[code]}), {disc_titles[code]} titles')
        for _d, w, c, dfp, p in scored[:14]:
            print(f"    ('{p}', {w}),   # n={c} df={dfp}")

    print('\n=== type -> discipline purity (support>=3) ===')
    td = defaultdict(Counter)
    for _t, typ, code in records:
        if typ and code:
            td[typ][code] += 1
    for typ in sorted(td):
        tot = sum(td[typ].values())
        top, n = td[typ].most_common(1)[0]
        if tot >= 3:
            flag = 'PURE' if n == tot else ''
            print(f"  {typ}: {top}({CODE2ID[top]}) {n}/{tot}={n/tot:.0%} {flag}"
                  f"  {dict(td[typ])}")


if __name__ == '__main__':
    main()
