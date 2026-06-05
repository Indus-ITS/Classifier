import openpyxl
from classifier.core.table_detect import (
    _header_span,
    MIN_HEADER_COLS,
    find_table,
    MIN_DATA_ROWS,
    file_has_table,
)


def test_header_span_finds_three_adjacent_text_cells():
    row = [None, "Sr No", "Area", "Line Number", None]
    start, width = _header_span(row)
    assert (start, width) == (1, 3)


def test_header_span_rejects_too_few_columns():
    row = [None, "Title", None, None]
    assert _header_span(row) is None


def test_header_span_ignores_pure_numbers_as_header():
    # a row of numbers is data, not a header
    row = [1, 2, 3, 4]
    assert _header_span(row) is None


def test_min_header_cols_is_three():
    assert MIN_HEADER_COLS == 3


def _mto_grid():
    header = [None, "Sr No", "Tag", "Description", "Qty"]
    rows = [[None, i, f"T-{i}", "VALVE", i * 2] for i in range(1, 9)]
    return [[None] * 5, header, *rows]


def test_find_table_accepts_real_table():
    hit = find_table(_mto_grid())
    assert hit is not None
    assert hit.n_data_rows >= MIN_DATA_ROWS
    assert hit.header_row == 1


def test_find_table_rejects_label_value_form():
    # CRS-style: each "data" row only fills 1-2 scattered cells
    grid = [
        ["DESCON TRANSMITTAL", "ADCO PDR", "DATE", "REVIEW CODE"],
        ["CTA-ED-SA-1576", None, None, None],
        ["REVIEWED BY", None, None, None],
        ["NAME", "Ahmed", None, None],
        ["POSITION", "CE", None, None],
    ]
    assert find_table(grid) is None


def test_find_table_rejects_short_table():
    header = ["A", "B", "C"]
    grid = [header, [1, 2, 3], [4, 5, 6]]  # only 2 data rows < MIN_DATA_ROWS
    assert find_table(grid) is None


def test_find_table_scans_past_leading_blank_rows():
    grid = [[None, None, None, None]] * 4 + _mto_grid()
    assert find_table(grid) is not None


def test_file_has_table_finds_table_behind_cover_sheet(tmp_path):
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["B2"] = "PROJECT TITLE"           # non-tabular cover
    idx = wb.create_sheet("Index")
    idx.append(["Sr No", "Area", "Line Number", "ISO Dwg"])
    for i in range(1, 9):
        idx.append([i, f"01{i:03d}P", f'3"-D-{i}', f"16-01-15-{i}"])
    p = tmp_path / "iso_index.xlsx"
    wb.save(p)

    hit = file_has_table(str(p))
    assert hit is not None
    assert hit[0] == "Index"          # sheet name
    assert hit[1].n_data_rows >= 5


def test_file_has_table_returns_none_for_form(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CRS"
    ws["B1"] = "ADCO PROJECT"
    ws["B2"] = "RESPONSE SHEET"
    ws["A4"] = "NAME"; ws["B4"] = "Ahmed"
    ws["A5"] = "POSITION"; ws["B5"] = "CE"
    p = tmp_path / "crs.xlsx"
    wb.save(p)

    assert file_has_table(str(p)) is None


def test_find_table_skips_blank_row_within_data():
    header = ["A", "B", "C"]
    data = [[1, 2, 3], [4, 5, 6], [7, 8, 9]]
    blank = [None, None, None]
    # blank row in the middle must NOT stop the count
    grid = [header, *data, blank, [10, 11, 12], [13, 14, 15]]
    hit = find_table(grid)
    assert hit is not None
    assert hit.n_data_rows == 5
