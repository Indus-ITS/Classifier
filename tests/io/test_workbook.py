import openpyxl
from classifier.io.workbook import file_has_table


def test_file_has_table_finds_table_behind_cover_sheet(tmp_path):
    wb = openpyxl.Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["B2"] = "PROJECT TITLE"
    idx = wb.create_sheet("Index")
    idx.append(["Sr No", "Area", "Line Number", "ISO Dwg"])
    for i in range(1, 9):
        idx.append([i, f"01{i:03d}P", f'3"-D-{i}', f"16-01-15-{i}"])
    p = tmp_path / "iso_index.xlsx"
    wb.save(p)

    hit = file_has_table(str(p))
    assert hit is not None
    assert hit[0] == "Index"
    assert hit[1].n_data_rows >= 5


def test_file_has_table_returns_none_for_form(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CRS"
    ws["B1"] = "ADCO PROJECT"
    ws["B2"] = "RESPONSE SHEET"
    ws["A4"] = "NAME"; ws["B4"] = "Ahmed"
    ws["A5"] = "POSITION"; ws["B5"] = "CE"
    p = tmp_path / "crs.xlsx"
    wb.save(p)

    assert file_has_table(str(p)) is None
